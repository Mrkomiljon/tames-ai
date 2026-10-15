#!/usr/bin/env python
"""Build the submission files from paper/*.md:  python paper/word/build_word.py
Needs pandoc (brew install pandoc) and python-docx (pip install python-docx).
Outputs paper/word/out/*.docx; equations become native Word equations, figures are embedded from paper/figures/.
"""
import subprocess
import sys
from pathlib import Path

PAPER = Path(__file__).resolve().parents[1]
HERE = PAPER / "word"
OUT = HERE / "out"
OUT.mkdir(exist_ok=True)
JOBS = [("manuscript_v3.md", "TAMES_AI_manuscript.docx"),
        ("supplementary.md", "TAMES_AI_supplementary.docx"),
        ("cover_letter_and_response.md", "TAMES_AI_cover_letter_and_response.docx")]


def shrink_tables(docx: Path, size_pt: float = 9):
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt
    d = Document(str(docx))
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.space_after = Pt(1); pf.space_before = Pt(1)
                    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    for r in p.runs:
                        r.font.size = Pt(size_pt)
    d.save(str(docx))


for src, dst in JOBS:
    cmd = ["pandoc", str(PAPER / src), "-o", str(OUT / dst),
           "-f", "markdown+tex_math_dollars+pipe_tables+fancy_lists+startnum+raw_html",
           "--reference-doc", str(HERE / "reference.docx"), "--resource-path", str(PAPER)]
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode:
        print(r.stderr); sys.exit(1)
    shrink_tables(OUT / dst)
    print("wrote", OUT / dst)
